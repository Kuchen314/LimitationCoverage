"""BLEU and BERTScore baselines, restricted to an exact paper-id list.

Same measurement as baseline_bleu_bertscore.py - both metrics in one pass, max
aggregation per REVIEWER2 (Gao et al. 2024 §5.1) - but the paper set is taken
from a file (e.g. coverage_table-FINAL.xlsx) so the baselines are computed on
exactly the same papers as similarity_coverage_final.py and the judge table.
A reconciliation report at the end names every requested paper that was not
scored, with the reason.

Two aggregations over the limitation list are reported, matching the other
instruments:
    mean - coverage-like: how much of the whole expert list is reflected
    max  - the REVIEWER2 statistic: how well the review matches its single
           best-matching expert concern
Each per-limitation score is itself a max over review units: a review need
only raise a concern somewhere.

Install: pip install bert-score supabase python-dotenv numpy scipy openpyxl tabulate
"""
import os
import json
import math
import re
import warnings
from collections import Counter

from supabase import create_client
from dotenv import load_dotenv
from tabulate import tabulate

warnings.filterwarnings("ignore", message=".*not writable.*")
warnings.filterwarnings("ignore", message=".*input array is constant.*")

load_dotenv()
supabase = create_client(os.environ["SUPABASE_URL"], os.environ["SUPABASE_SERVICE_KEY"])

# ------------------------------- CONFIG -------------------------------
TABLE = "paper"
REVIEW_COLS = {1: "review3", 2: "review_edited3"}
LIMITATIONS_COL = "limitations"
JUDGE_COL = "comparison3"

# --- the exact paper set ---------------------------------------------
RESTRICT_TO_IDS = "coverage_table-FINAL.xlsx"   # None = whole table
ID_COLUMN = "Paper ID"        # column IN THE FILE holding the identifier
DB_ID_COL = "id"              # column IN THE DATABASE those values match.
                              # Set to None to instead read the id out of a
                              # JSON column (e.g. {"name": "316.pdf"}), using:
ID_JSON_COL = None            # e.g. "parsed_pdf"
ID_JSON_KEY = None            # e.g. "name"  (None = the column is plain text)
ID_JSON_REGEX = r"(\d+)"

IGNORE_COL = "ignore"
SKIP_IGNORED = True           # ignored when RESTRICT_TO_IDS is set: an explicit
                              # id list is authoritative, so flagged papers still
                              # appear and are reported in the reconciliation
SKIP_ON_MISMATCH = True       # drop papers whose limitation count != judge total
MIN_LIMITATIONS = 3

BERT_MODEL = "roberta-large"
RESCALE = True
WINDOW = 2                    # also match against windows of N sentences
CRITICAL_ONLY = True          # only the criticism sections of the review

OUTPUT_XLSX = "baseline_final.xlsx"
OUTPUT_JSON = "baseline_final_scores.json"
# ----------------------------------------------------------------------


# ------------------------------ parsing ------------------------------

def extract_text(val):
    if val is None:
        return None
    if isinstance(val, dict):
        return val.get("review_outline") or val.get("review")
    if isinstance(val, str):
        s = val.strip()
        if s.startswith("{"):
            try:
                return extract_text(json.loads(s))
            except json.JSONDecodeError:
                pass
        return s
    return None


def split_limitations(val):
    if val is None:
        return []
    if isinstance(val, list):
        return [str(x).strip() for x in val if str(x).strip()]
    if isinstance(val, dict):
        for v in val.values():
            if isinstance(v, list):
                return [str(x).strip() for x in v if str(x).strip()]
        val = json.dumps(val)
    if isinstance(val, str):
        s = val.strip()
        if s.startswith("["):
            try:
                return [str(x).strip() for x in json.loads(s)]
            except json.JSONDecodeError:
                pass
        items = [re.sub(r"^[\s\*\-•]*(\d+[\.\)])?\s*", "", line).strip()
                 for line in s.split("\n")]
        return [x for x in items if len(x) > 10]
    return []


def critical_sections(text):
    """Sections 2 and 3 of the numbered review carry the criticism; the summary
    shares vocabulary with every limitation without expressing a concern."""
    parts = re.split(r"\n(?=\d+\.\s)", "\n" + text)
    keep = [p for p in parts if re.match(r"\s*[23]\.\s", p)]
    return "\n".join(keep) if keep else text


def make_units(text, window=WINDOW):
    sents = [p.strip() for p in re.split(r"(?<=[.!?])\s+|\n+", text)
             if len(p.strip()) > 15]
    units = list(sents)
    for w in range(2, window + 1):
        for i in range(len(sents) - w + 1):
            units.append(" ".join(sents[i:i + w]))
    return units


def judge_coverage(comparison, review_idx):
    if isinstance(comparison, str):
        try:
            comparison = json.loads(comparison)
        except json.JSONDecodeError:
            return None, None
    if not isinstance(comparison, dict):
        return None, None
    m = re.search(r"(\d+)/(\d+)", str(comparison.get(f"Review {review_idx} Coverage", "")))
    return (int(m.group(1)), int(m.group(2))) if m else (None, None)


# ------------------------------- BLEU -------------------------------

def _tokens(s):
    return re.findall(r"[a-z0-9']+", s.lower())


def _ngrams(toks, n):
    return Counter(tuple(toks[i:i + n]) for i in range(len(toks) - n + 1))


def sentence_bleu(candidate, reference, max_n=4):
    """Smoothed sentence BLEU: clipped n-gram precision for n=1..4, geometric
    mean, add-one smoothing (so one missing order does not zero short texts),
    brevity penalty."""
    c, r = _tokens(candidate), _tokens(reference)
    if not c or not r:
        return 0.0
    log_sum = 0.0
    for n in range(1, max_n + 1):
        cn, rn = _ngrams(c, n), _ngrams(r, n)
        match = sum(min(cnt, rn[g]) for g, cnt in cn.items())
        total = max(sum(cn.values()), 1)
        log_sum += math.log((match + 1) / (total + 1))
    bp = 1.0 if len(c) > len(r) else math.exp(1 - len(r) / max(len(c), 1))
    return bp * math.exp(log_sum / max_n)


def bleu_max_per_limitation(limitations, units):
    return [max((sentence_bleu(u, lim) for u in units), default=0.0)
            for lim in limitations]


# ----------------------------- BERTScore -----------------------------

def bertscore_max_per_limitation(scorer, limitations, units):
    """All (unit, limitation) pairs scored in ONE batch, reshaped
    limitation-major, reduced with max over the unit axis."""
    cands, refs = [], []
    for lim in limitations:
        for u in units:
            cands.append(u)
            refs.append(lim)
    if not cands:
        return [0.0] * len(limitations)
    _, _, f1 = scorer.score(cands, refs)
    f1 = f1.tolist()
    nu = len(units)
    return [max(f1[i * nu:(i + 1) * nu]) for i in range(len(limitations))]


# --------------------------- id resolution ---------------------------

def _idkey(v):
    """374, 374.0 and ' 374 ' all compare equal - guards against the
    int-in-spreadsheet vs text-in-database mismatch."""
    s = str(v).strip()
    try:
        return str(int(float(s)))
    except ValueError:
        return s.casefold()


def load_id_filter():
    if not RESTRICT_TO_IDS:
        return None
    ids = []
    if str(RESTRICT_TO_IDS).lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        rows = list(load_workbook(RESTRICT_TO_IDS, read_only=True).active
                    .iter_rows(values_only=True))
    else:
        import csv
        with open(RESTRICT_TO_IDS, newline="", encoding="utf-8") as f:
            rows = list(csv.reader(f))
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    col = header.index(ID_COLUMN) if ID_COLUMN in header else 0
    for r in rows[1:]:
        if col >= len(r) or r[col] is None or not str(r[col]).strip():
            continue
        s = str(r[col]).strip()
        if s.upper() in ("AVERAGE", "MEAN", "MEDIAN", "TOTAL", "SUM"):
            continue
        try:
            ids.append(int(float(s)))
        except ValueError:
            ids.append(s)
    out, seen = [], set()
    for i in ids:
        k = _idkey(i)
        if k not in seen:
            seen.add(k)
            out.append(i)
    target = f"{TABLE}.{DB_ID_COL}" if DB_ID_COL else f"{TABLE}.{ID_JSON_COL}"
    print(f"id filter: {len(out)} unique ids from {RESTRICT_TO_IDS} "
          f"(column '{ID_COLUMN}') -> matched against {target}")
    print(f"  first 5: {out[:5]}")
    return out


def build_paper_id_map():
    """{paper id key -> db key}, {db key -> display id}, when the identifier
    lives inside a JSON column rather than being a column of its own."""
    if ID_JSON_KEY is None:
        rows = supabase.table(TABLE).select(f"id, {ID_JSON_COL}").execute().data
        get = lambda r: r.get(ID_JSON_COL)
    else:
        try:
            rows = (supabase.table(TABLE)
                    .select(f"id, _pid:{ID_JSON_COL}->>{ID_JSON_KEY}").execute()).data
            get = lambda r: r.get("_pid")
        except Exception:
            rows = supabase.table(TABLE).select(f"id, {ID_JSON_COL}").execute().data
            def get(r):
                v = r.get(ID_JSON_COL)
                if isinstance(v, str):
                    try:
                        v = json.loads(v)
                    except json.JSONDecodeError:
                        return v
                return v.get(ID_JSON_KEY) if isinstance(v, dict) else v
    mapping, display = {}, {}
    for r in rows:
        raw = get(r)
        if raw is None:
            continue
        m = re.search(ID_JSON_REGEX, str(raw))
        if not m or _idkey(m.group(1)) in mapping:
            continue
        mapping[_idkey(m.group(1))] = r["id"]
        try:
            display[r["id"]] = int(m.group(1))
        except ValueError:
            display[r["id"]] = m.group(1)
    print(f"id map: {len(mapping)} paper ids resolved from {ID_JSON_COL}")
    return mapping, display


def fetch_rows(id_filter):
    """(rows, dbid_to_paperid). The second value is None when the id is a
    plain column, since no translation is needed."""
    cols = ["id", LIMITATIONS_COL] + list(REVIEW_COLS.values())
    if JUDGE_COL:
        cols.append(JUDGE_COL)
    if DB_ID_COL and DB_ID_COL != "id":
        cols.append(DB_ID_COL)

    def base():
        q = (supabase.table(TABLE)
             .select(", ".join(dict.fromkeys(cols)))
             .not_.is_(REVIEW_COLS[1], "null")
             .not_.is_(LIMITATIONS_COL, "null"))
        if IGNORE_COL and SKIP_IGNORED and id_filter is None:
            q = q.not_.is_(IGNORE_COL, "true")
        return q

    if id_filter is None:
        rows = base().execute().data
        print(f"fetched {len(rows)} rows (no id restriction)")
        return rows, None

    if DB_ID_COL:
        rows, CHUNK = [], 100
        for i in range(0, len(id_filter), CHUNK):
            rows.extend(base().in_(DB_ID_COL, id_filter[i:i + CHUNK]).execute().data)
        print(f"fetched {len(rows)} of {len(id_filter)} requested rows")
        return rows, None

    id_map, back = build_paper_id_map()
    wanted = [id_map[_idkey(p)] for p in id_filter if _idkey(p) in id_map]
    rows, CHUNK = [], 100
    for i in range(0, len(wanted), CHUNK):
        rows.extend(base().in_("id", wanted[i:i + CHUNK]).execute().data)
    print(f"fetched {len(rows)} of {len(id_filter)} requested rows")
    return rows, back


# ------------------------------- main -------------------------------

def main(max_papers=None):
    from bert_score import BERTScorer
    scorer = BERTScorer(model_type=BERT_MODEL, lang="en", rescale_with_baseline=RESCALE)

    id_filter = load_id_filter()
    rows, dbid_to_pid = fetch_rows(id_filter)

    def paper_id_of(row):
        if dbid_to_pid is not None:
            return dbid_to_pid.get(row["id"], f"db:{row['id']}")
        return row[DB_ID_COL] if DB_ID_COL else row["id"]

    truncated = 0
    if max_papers and max_papers < len(rows):
        truncated = len(rows) - max_papers
        rows = rows[:max_papers]
        print(f"NOTE: max_papers={max_papers} - {truncated} rows dropped; the "
              f"reconciliation will list them as missing.")

    prepared, dropped = [], {}
    for k, row in enumerate(rows):
        pid = paper_id_of(row)
        limitations = split_limitations(row[LIMITATIONS_COL])
        if len(limitations) < MIN_LIMITATIONS:
            dropped[pid] = f"only {len(limitations)} limitations"
            continue
        jc1, jt1 = judge_coverage(row.get(JUDGE_COL), 1) if JUDGE_COL else (None, None)
        jc2, _ = judge_coverage(row.get(JUDGE_COL), 2) if JUDGE_COL else (None, None)
        if SKIP_ON_MISMATCH and jt1 is not None and jt1 != len(limitations):
            dropped[pid] = f"denominator mismatch ({len(limitations)} vs {jt1})"
            continue
        rec = {"id": pid, "total": len(limitations), "limitations": limitations,
               "judge_cov1": jc1, "judge_cov2": jc2}
        ok = True
        for idx, col in REVIEW_COLS.items():
            text = extract_text(row.get(col))
            if not text:
                ok = False
                break
            rec[f"words_full{idx}"] = len(_tokens(text))
            if CRITICAL_ONLY:
                text = critical_sections(text)
            units = make_units(text)
            if not units:
                ok = False
                break
            rec[f"units{idx}"] = units
            rec[f"words{idx}"] = len(_tokens(text))
            toks = _tokens(text)
            rec[f"ttr{idx}"] = len(set(toks)) / len(toks) if toks else 0.0
        if ok:
            prepared.append(rec)
        else:
            dropped[pid] = "missing or empty review text"

    if id_filter is not None:
        got = {_idkey(p["id"]) for p in prepared}
        by_key = {_idkey(k): v for k, v in dropped.items()}
        missing = [i for i in id_filter if _idkey(i) not in got]
        print("\n=== id reconciliation ===")
        print(f"requested : {len(id_filter)}")
        print(f"scored    : {len(prepared)}")
        print(f"missing   : {len(missing)}"
              + (f"  ({truncated} only because max_papers was set)" if truncated else ""))
        for i in missing:
            print(f"  {i}: {by_key.get(_idkey(i), 'no row returned from the database')}")
        if not missing:
            print("  -> exact match with the requested paper set")

    print(f"\nscoring {len(prepared)} papers\n")
    for i, p in enumerate(prepared):
        for idx in REVIEW_COLS:
            units = p[f"units{idx}"]
            bleu = bleu_max_per_limitation(p["limitations"], units)
            bert = bertscore_max_per_limitation(scorer, p["limitations"], units)
            p[f"bleu_scores{idx}"] = [round(x, 5) for x in bleu]
            p[f"bert_scores{idx}"] = [round(x, 5) for x in bert]
            p[f"bleu_mean{idx}"] = sum(bleu) / len(bleu)
            p[f"bleu_max{idx}"] = max(bleu)
            p[f"bert_mean{idx}"] = sum(bert) / len(bert)
            p[f"bert_max{idx}"] = max(bert)
        p.pop("units1", None)
        p.pop("units2", None)
        print(f"[{i+1}/{len(prepared)}] paper {p['id']}: "
              f"BLEU {p['bleu_mean1']:.3f}/{p['bleu_mean2']:.3f}  "
              f"BERT {p['bert_mean1']:.3f}/{p['bert_mean2']:.3f}")

    # ------------------------------ table ------------------------------
    table = [[p["id"], p["total"], p["words1"], p["words2"],
              f"{p['bleu_mean1']:.3f}", f"{p['bleu_mean2']:.3f}",
              f"{p['bleu_max1']:.3f}", f"{p['bleu_max2']:.3f}",
              f"{p['bert_mean1']:.3f}", f"{p['bert_mean2']:.3f}",
              f"{p['bert_max1']:.3f}", f"{p['bert_max2']:.3f}",
              f"{p['judge_cov1']}/{p['total']}" if p["judge_cov1"] is not None else "—",
              f"{p['judge_cov2']}/{p['total']}" if p["judge_cov2"] is not None else "—"]
             for p in prepared]
    print()
    print(tabulate(table, headers=["Paper", "Total", "w R1", "w R2",
                                   "BLEU mean R1", "BLEU mean R2",
                                   "BLEU max R1", "BLEU max R2",
                                   "BERT mean R1", "BERT mean R2",
                                   "BERT max R1", "BERT max R2",
                                   "Judge R1", "Judge R2"],
                   tablefmt="rounded_outline"))

    # ---------------------------- statistics ----------------------------
    try:
        import numpy as np
        from scipy import stats

        print("\n=== condition 1 vs condition 2 (paired) ===")
        for metric in ("bleu", "bert"):
            for agg in ("mean", "max"):
                a = np.array([p[f"{metric}_{agg}1"] for p in prepared])
                b = np.array([p[f"{metric}_{agg}2"] for p in prepared])
                d = a - b
                try:
                    pv = stats.wilcoxon(a, b).pvalue
                except ValueError:
                    pv = float("nan")
                sd = d.std(ddof=1)
                print(f"{metric.upper():5s} {agg:4s}: R1={a.mean():.4f} R2={b.mean():.4f} "
                      f"Δ={d.mean():+.4f}  wilcoxon p={pv:.3g}  "
                      f"d={d.mean()/sd if sd else float('nan'):.2f}")

        if JUDGE_COL:
            print("\n=== correlation with judge coverage (Spearman) ===")
            for metric in ("bleu", "bert"):
                for agg in ("mean", "max"):
                    for idx in (1, 2):
                        pr = [(p[f"{metric}_{agg}{idx}"], p[f"judge_cov{idx}"] / p["total"])
                              for p in prepared if p[f"judge_cov{idx}"] is not None]
                        if len(pr) > 5:
                            x, y = zip(*pr)
                            r = stats.spearmanr(x, y)
                            print(f"{metric.upper():5s} {agg:4s} R{idx}: "
                                  f"rho={r.statistic:+.3f} (p={r.pvalue:.3g}, n={len(pr)})")

        print("\n=== verbosity diagnostic ===")
        w1 = np.array([p["words1"] for p in prepared], float)
        w2 = np.array([p["words2"] for p in prepared], float)
        print(f"words (scored text): R1 {w1.mean():.0f}  R2 {w2.mean():.0f}  "
              f"Δ={w1.mean()-w2.mean():+.0f}  wilcoxon p={stats.wilcoxon(w1, w2).pvalue:.3g}")
        t1 = np.array([p["ttr1"] for p in prepared], float)
        t2 = np.array([p["ttr2"] for p in prepared], float)
        print(f"type-token ratio   : R1 {t1.mean():.3f}  R2 {t2.mean():.3f}  "
              f"wilcoxon p={stats.wilcoxon(t1, t2).pvalue:.3g}")
        words = np.concatenate([w1, w2])
        print(f"\nmetric vs word count (reviews pooled, n={len(words)}):")
        for metric in ("bleu", "bert"):
            for agg in ("mean", "max"):
                v = np.concatenate([[p[f"{metric}_{agg}1"] for p in prepared],
                                    [p[f"{metric}_{agg}2"] for p in prepared]])
                r = stats.spearmanr(words, v)
                print(f"  {metric.upper():5s} {agg:4s}: rho={r.statistic:+.3f} (p={r.pvalue:.3g})")
        jp = [(p["words1"], p["judge_cov1"] / p["total"]) for p in prepared
              if p["judge_cov1"] is not None]
        jp += [(p["words2"], p["judge_cov2"] / p["total"]) for p in prepared
               if p["judge_cov2"] is not None]
        if len(jp) > 5:
            jw, jc = map(np.array, zip(*jp))
            r = stats.spearmanr(jw, jc)
            print(f"  JUDGE      : rho={r.statistic:+.3f} (p={r.pvalue:.3g}, n={len(jw)})")
        print("\nper-paper differences (R1 - R2):")
        dw = w1 - w2
        for metric in ("bleu", "bert"):
            for agg in ("mean", "max"):
                dm = np.array([p[f"{metric}_{agg}1"] - p[f"{metric}_{agg}2"]
                               for p in prepared])
                r = stats.spearmanr(dw, dm)
                print(f"  Δwords vs Δ{metric.upper()}_{agg}: rho={r.statistic:+.3f} "
                      f"(p={r.pvalue:.3g})")
        dj = [(p["words1"] - p["words2"],
               (p["judge_cov1"] - p["judge_cov2"]) / p["total"]) for p in prepared
              if p["judge_cov1"] is not None and p["judge_cov2"] is not None]
        if len(dj) > 5:
            a, b = map(np.array, zip(*dj))
            r = stats.spearmanr(a, b)
            print(f"  Δwords vs Δjudge      : rho={r.statistic:+.3f} "
                  f"(p={r.pvalue:.3g}, n={len(a)})")
    except ImportError:
        print("scipy/numpy not installed - skipping statistics")

    # ------------------------------ output ------------------------------
    with open(OUTPUT_JSON, "w", encoding="utf-8") as f:
        json.dump(prepared, f, indent=1)
    print(f"\nraw per-limitation scores -> {OUTPUT_JSON}")

    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "baselines"
    ws.append(["Paper ID", "Total", "words R1", "words R2", "TTR R1", "TTR R2",
               "BLEU mean R1", "BLEU mean R2", "BLEU max R1", "BLEU max R2",
               "BERT mean R1", "BERT mean R2", "BERT max R1", "BERT max R2",
               "Judge cov R1", "Judge cov R2"])
    for p in prepared:
        ws.append([p["id"], p["total"], p["words1"], p["words2"],
                   round(p["ttr1"], 4), round(p["ttr2"], 4),
                   round(p["bleu_mean1"], 5), round(p["bleu_mean2"], 5),
                   round(p["bleu_max1"], 5), round(p["bleu_max2"], 5),
                   round(p["bert_mean1"], 5), round(p["bert_mean2"], 5),
                   round(p["bert_max1"], 5), round(p["bert_max2"], 5),
                   p["judge_cov1"], p["judge_cov2"]])
    wb.save(OUTPUT_XLSX)
    print(f"Saved to {OUTPUT_XLSX}")


if __name__ == "__main__":
    main()   # add max_papers=20 for a quick smoke test
