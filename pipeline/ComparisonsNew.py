import os
import json
import re
from supabase import create_client
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from tabulate import tabulate

load_dotenv()

supabase = create_client(os.environ["SUPABASE_URL"], os.environ["SUPABASE_SERVICE_KEY"])

SCORE_FIELDS = ["Soundness", "Presentation", "Significance", "Originality"]


def parse_coverage(coverage_val):
    # coverage_val may be a string like "5/9 limitations covered..."
    # or a dict like {"covered": 5, "total": 9, ...} depending on how the model returned it
    if isinstance(coverage_val, dict):
        covered = coverage_val.get("covered") or coverage_val.get("count") or coverage_val.get("covered_count")
        total = coverage_val.get("total") or coverage_val.get("total_count")
        # some responses store the actual items instead of a count
        if isinstance(covered, list):
            covered = len(covered)
        if isinstance(total, list):
            total = len(total)
        if covered is not None and total:
            return int(covered), int(total)
        # fall back to searching string representation of the dict
        coverage_val = str(coverage_val)

    if not isinstance(coverage_val, str):
        return None, None

    match = re.search(r"(\d+)/(\d+)", coverage_val)
    if match:
        return int(match.group(1)), int(match.group(2))
    return None, None


def _to_int(val):
    # Scores should be bare ints, but tolerate "2", "2: fair" or "2 (fair)".
    # bool is excluded because in Python True/False are subclasses of int.
    if isinstance(val, bool):
        return None
    if isinstance(val, int):
        return val
    if isinstance(val, str):
        m = re.match(r"\s*(\d+)", val)
        if m:
            return int(m.group(1))
    return None


def parse_scores(comparison):
    """Parse the "Scores" object of the comparison JSON.

    Expected shape:
      {"Scores": {"Review 1": {"Soundness": 3, ...}, "Review 2": {...}, "Disagreements": "..."}}

    Returns {1: {field: int|None}, 2: {...}}.
    """
    result = {1: {f: None for f in SCORE_FIELDS}, 2: {f: None for f in SCORE_FIELDS}}
    scores = comparison.get("Scores") if isinstance(comparison, dict) else None
    if not isinstance(scores, dict):
        return result

    for idx in (1, 2):
        # tolerate small key variations: "Review 1", "review 1", "Review1", "review_1"
        review = None
        for key, val in scores.items():
            if isinstance(val, dict) and re.fullmatch(rf"review[\s_]?{idx}", key, flags=re.IGNORECASE):
                review = val
                break
        if review is None:
            continue
        for field in SCORE_FIELDS:
            # case-insensitive field lookup
            for k, v in review.items():
                if k.lower() == field.lower():
                    result[idx][field] = _to_int(v)
                    break
    return result


def parse_decision(comparison):
    """Parse the "Decision" string of the comparison JSON.

    Expected shape: "Review 1: score 4, Accept; Review 2: score 6, Accept. ..."

    Returns {1: (score|None, accepted|None), 2: (...)}.
    """
    result = {1: (None, None), 2: (None, None)}
    decision = comparison.get("Decision") if isinstance(comparison, dict) else None
    if not isinstance(decision, str):
        decision = json.dumps(decision) if decision is not None else ""

    for idx in (1, 2):
        # "Review 1: score 4, Accept" — allow small punctuation/wording drift:
        # \D{0,15} lets e.g. "Review 1 gives score 4" through, and the
        # accept/reject word must follow the number within a few characters.
        m = re.search(
            rf"review\s*{idx}\D{{0,15}}score\D{{0,3}}(\d+)\W{{0,5}}(accept|reject)",
            decision,
            flags=re.IGNORECASE,
        )
        if m:
            result[idx] = (int(m.group(1)), m.group(2).lower() == "accept")
    return result


def column_name(option):
    # option 1 -> "comparison" (the first column has no number suffix),
    # option 2 -> "comparison2", option 3 -> "comparison3", ...
    return "comparison" if option == 1 else f"comparison{option}"


def fetch_comparisons(option=1, skip_ignored=False):
    column = column_name(option)
    query = (
        supabase.table("paper")
        .select(f"id, {column}")
        .not_.is_(column, "null")
    )
    if skip_ignored:
        # keep papers where ignore is false OR null (never explicitly flagged);
        # only ignore=true rows are excluded
        query = query.not_.is_("ignore", "true")
    response = query.execute()
    return response.data


def build_table(option = 1, output_path="coverage_table2.xlsx", skip_ignored=False,
                min_total=None):
    # min_total: if set, papers whose coverage denominator (the total number
    # of expert limitations, e.g. the 3 in "2/3") is below this value are
    # filtered out. min_total=3 keeps totals of 3 and up.
    rows = fetch_comparisons(option, skip_ignored)
    column = column_name(option)

    wb = Workbook()
    ws = wb.active
    ws.title = "Coverage"

    # header styling
    header_font = Font(bold=True, color="FFFFFF", name="Arial")
    header_fill = PatternFill("solid", start_color="2F5496")
    header_align = Alignment(horizontal="center", vertical="center")

    # per review: coverage, the 4 scores, recommendation score, decision
    per_review_cols = ["Coverage", "Sound", "Pres", "Signif", "Orig", "Rec", "Decision"]
    # difference columns (R1 - R2), computed per paper
    diff_cols = ["Δ Coverage", "Δ Sound", "Δ Pres", "Δ Signif", "Δ Orig", "Δ Rec", "Agree?"]
    headers = ["Paper ID"] \
        + [f"R1 {c}" for c in per_review_cols] \
        + [f"R2 {c}" for c in per_review_cols] \
        + diff_cols
    n_cols = len(headers)

    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    # accumulators for the average row, kept separately per review:
    # coverage ratios, one list per score field, recommendation scores,
    # and accepted booleans (averaged into an accept rate)
    stats = {
        idx: {"coverage": [], "rec": [], "accepted": [], **{f: [] for f in SCORE_FIELDS}}
        for idx in (1, 2)
    }
    # accumulators for the R1 - R2 differences and decision agreement
    diff_stats = {"coverage": [], "rec": [], "agree": [], **{f: [] for f in SCORE_FIELDS}}
    table_data = []  # shared list used for both Excel and console output
    data_row = 2

    for row in rows:
        paper_id = row["id"]
        comparison = row[column]

        if isinstance(comparison, str):
            try:
                comparison = json.loads(comparison)
            except json.JSONDecodeError:
                print(f"Skipping paper {paper_id}: could not parse comparison JSON")
                continue

        r1_str = comparison.get("Review 1 Coverage", "")
        r2_str = comparison.get("Review 2 Coverage", "")

        r1_covered, r1_total = parse_coverage(r1_str)
        r2_covered, r2_total = parse_coverage(r2_str)

        if r1_covered is None or r2_covered is None:
            print(f"Skipping paper {paper_id}: coverage not in expected format")
            continue

        # denominator filter: skip papers with too few expert limitations
        if min_total is not None and (r1_total < min_total or r2_total < min_total):
            print(f"Skipping paper {paper_id}: total {min(r1_total, r2_total)} < min_total {min_total}")
            continue

        coverage = {1: (r1_covered, r1_total), 2: (r2_covered, r2_total)}
        scores = parse_scores(comparison)
        decision = parse_decision(comparison)

        # assemble the full row: paper id + 7 columns per review
        excel_row = [paper_id]
        for idx in (1, 2):
            covered, total = coverage[idx]
            excel_row.append(f"{covered}/{total}")
            stats[idx]["coverage"].append(covered / total)

            for field in SCORE_FIELDS:
                val = scores[idx][field]
                excel_row.append(val if val is not None else "—")
                if val is not None:
                    stats[idx][field].append(val)

            rec, accepted = decision[idx]
            excel_row.append(rec if rec is not None else "—")
            if rec is not None:
                stats[idx]["rec"].append(rec)
            excel_row.append("Accept" if accepted else "Reject" if accepted is not None else "—")
            if accepted is not None:
                stats[idx]["accepted"].append(accepted)

        # difference columns: R1 - R2 (positive = Review 1 higher)
        def _diff(a, b, store, fmt="{:+d}"):
            if a is None or b is None:
                return "—"
            d = a - b
            store.append(d)
            return fmt.format(d)

        c1, t1 = coverage[1]
        c2, t2 = coverage[2]
        excel_row.append(_diff(c1 / t1, c2 / t2, diff_stats["coverage"], fmt="{:+.1%}"))
        for field in SCORE_FIELDS:
            excel_row.append(_diff(scores[1][field], scores[2][field], diff_stats[field]))
        excel_row.append(_diff(decision[1][0], decision[2][0], diff_stats["rec"]))
        acc1, acc2 = decision[1][1], decision[2][1]
        if acc1 is None or acc2 is None:
            excel_row.append("—")
        else:
            agree = acc1 == acc2
            diff_stats["agree"].append(agree)
            excel_row.append("Yes" if agree else "No")

        table_data.append(excel_row)

        for col, value in enumerate(excel_row, start=1):
            cell = ws.cell(row=data_row, column=col, value=value)
            if col > 1:
                cell.alignment = Alignment(horizontal="center")

        if data_row % 2 == 0:
            row_fill = PatternFill("solid", start_color="DCE6F1")
            for col in range(1, n_cols + 1):
                ws.cell(row=data_row, column=col).fill = row_fill

        data_row += 1

    # average row: coverage %, mean of each score, mean recommendation, accept rate %
    def _avg(vals, pct=False, signed=False):
        if not vals:
            return "—"
        mean = sum(vals) / len(vals)
        spec = ("+" if signed else "") + (".1%" if pct else ".2f")
        return format(mean, spec)

    avg_row = ["AVERAGE"]
    for idx in (1, 2):
        avg_row.append(_avg(stats[idx]["coverage"], pct=True))
        for field in SCORE_FIELDS:
            avg_row.append(_avg(stats[idx][field]))
        avg_row.append(_avg(stats[idx]["rec"]))
        avg_row.append(_avg([int(a) for a in stats[idx]["accepted"]], pct=True))  # accept rate

    # averages of the difference columns; "Agree?" becomes the agreement rate
    avg_row.append(_avg(diff_stats["coverage"], pct=True, signed=True))
    for field in SCORE_FIELDS:
        avg_row.append(_avg(diff_stats[field], signed=True))
    avg_row.append(_avg(diff_stats["rec"], signed=True))
    avg_row.append(_avg([int(a) for a in diff_stats["agree"]], pct=True))  # agreement rate

    table_data.append(avg_row)

    avg_fill = PatternFill("solid", start_color="F2F2F2")
    avg_font = Font(bold=True, name="Arial")

    for col, value in enumerate(avg_row, start=1):
        cell = ws.cell(row=data_row, column=col, value=value)
        cell.font = avg_font
        cell.fill = avg_fill
        cell.alignment = Alignment(horizontal="center")

    ws.column_dimensions["A"].width = 12
    for col in range(2, n_cols + 1):
        ws.column_dimensions[get_column_letter(col)].width = 11

    # print to console
    print(tabulate(
        table_data,
        headers=headers,
        tablefmt="rounded_outline",
        colalign=["left"] + ["center"] * (n_cols - 1),
    ))

    wb.save(output_path)
    print(f"Saved to {output_path}")


if __name__ == "__main__":
    build_table(option= 3,output_path="coverage_table-FINAL.xlsx", skip_ignored=True, min_total=3 )
